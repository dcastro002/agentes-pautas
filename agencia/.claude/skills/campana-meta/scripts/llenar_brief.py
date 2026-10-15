#!/usr/bin/env python3
"""Llena la plantilla de brief de pauta (Excel) a partir de un JSON.

Uso:
  python3 llenar_brief.py <datos.json> <plantilla.xlsx> <salida.xlsx>

- Respeta el formato de la plantilla: una pestaña, datos en B1-B8,
  anuncios en las filas 11-18 y datos para el agente en K2-K13.
- Valida cada campo contra las opciones que acepta el agente de pautas.
- Cuenta, página y píxel (K11-K13) vienen del JSON (tomados de Cuentas.md)
  y se escriben como valores.
- Si algo no es válido, no escribe nada y sale con código 1.
"""
import json, sys, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: pip3 install --user openpyxl")

CELDAS = {
    "cliente": "B1", "nombre": "B2", "plataforma": "B3", "objetivo": "B4",
    "geo": "B5", "inicio": "B6", "fin": "B7", "presupuesto": "B8",
    "tipo_presupuesto": "K2", "evento": "K3", "categoria": "K4",
    "edad_min": "K5", "edad_max": "K6", "genero": "K7", "intereses": "K8",
    "ubicaciones": "K9", "instagram": "K10",
    "cuenta": "K11", "pagina": "K12", "pixel": "K13",
}
COL_ANUNCIO = {"copy": "A", "titulo": "B", "cta": "C", "url": "D",
               "formato": "G", "archivo": "H"}
FILAS = range(11, 19)  # máximo 8 anuncios

OPCIONES = {
    "plataforma": ["Meta", "Instagram", "Facebook"],
    "objetivo": ["Conversiones", "Tráfico", "Awareness", "Ventas por Catálogo",
                 "Interacción", "Alcance"],
    "tipo_presupuesto": ["Diario", "Total"],
    "evento": ["Lead", "Contact", "CompleteRegistration", "Schedule",
               "SubmitApplication", "Purchase", "ViewContent", "No aplica"],
    "categoria": ["Ninguna", "Vivienda", "Empleo", "Crédito", "Temas sociales o política"],
    "genero": ["Todos", "Mujeres", "Hombres"],
    "ubicaciones": ["Advantage+ (automáticas)", "Facebook e Instagram",
                    "Solo Facebook", "Solo Instagram"],
    "cta": ["Más información", "Comprar ahora", "Registrarte", "Contactarnos",
            "Enviar mensaje", "Descargar", "Reservar", "Suscribirte",
            "Solicitar hora", "Ver más"],
    "formato": ["Imagen única"],
}
OBLIGATORIOS = ["cliente", "nombre", "plataforma", "objetivo", "geo", "inicio", "fin",
                "presupuesto", "tipo_presupuesto", "categoria", "cuenta", "pagina", "pixel"]


def main(datos_p, plantilla_p, salida_p):
    d = json.load(open(datos_p, encoding="utf-8"))
    errores = []

    def opcion(campo, valor, donde):
        if campo in OPCIONES and valor not in OPCIONES[campo]:
            errores.append(f"{donde} = {valor!r} no es válido. Opciones: {OPCIONES[campo]}")

    for k in OBLIGATORIOS:
        if d.get(k) in (None, "", 0):
            errores.append(f"Falta {k} ({CELDAS[k]})")
    for k, celda in CELDAS.items():
        if d.get(k) not in (None, ""):
            opcion(k, d[k], f"{k} ({celda})")
    if d.get("objetivo") == "Conversiones" and d.get("evento") in (None, "", "No aplica"):
        errores.append("El objetivo es Conversiones: falta el evento (K3)")
    for k in ("inicio", "fin"):
        try:
            if d.get(k):
                datetime.datetime.strptime(d[k], "%Y-%m-%d")
        except ValueError:
            errores.append(f"{k} debe tener formato AAAA-MM-DD")
    if d.get("inicio") and d.get("fin") and d["fin"] <= d["inicio"]:
        errores.append("La fecha de fin debe ser posterior a la de inicio")

    anuncios = d.get("anuncios") or []
    if not anuncios:
        errores.append("Falta al menos un anuncio")
    if len(anuncios) > len(FILAS):
        errores.append(f"Máximo {len(FILAS)} anuncios")
    for i, a in enumerate(anuncios):
        fila = FILAS[0] + i
        for k, col in COL_ANUNCIO.items():
            if not a.get(k):
                errores.append(f"Anuncio {i+1}: falta {k} ({col}{fila})")
        if a.get("cta"):
            opcion("cta", a["cta"], f"Anuncio {i+1} cta (C{fila})")
        if a.get("formato"):
            opcion("formato", a["formato"], f"Anuncio {i+1} formato (G{fila})")
        if a.get("url") and not str(a["url"]).startswith("https://"):
            errores.append(f"Anuncio {i+1}: la URL debe empezar con https://")

    if errores:
        print("NO SE GENERÓ EL BRIEF:")
        for e in errores:
            print(" -", e)
        sys.exit(1)

    wb = openpyxl.load_workbook(plantilla_p)
    ws = wb.worksheets[0]
    for celda in CELDAS.values():
        ws[celda].value = None
    for fila in FILAS:
        for col in "ABCDEFGH":
            ws[f"{col}{fila}"].value = None

    for k, celda in CELDAS.items():
        v = d.get(k)
        if v in (None, ""):
            continue
        if k in ("inicio", "fin"):
            v = datetime.datetime.strptime(v, "%Y-%m-%d")
            ws[celda].number_format = "yyyy-mm-dd"
        if k == "cuenta":
            v = str(v)
        ws[celda].value = v
    for i, a in enumerate(anuncios):
        fila = FILAS[0] + i
        for k, col in COL_ANUNCIO.items():
            ws[f"{col}{fila}"].value = a[k]

    wb.save(salida_p)

    # Verificar leyendo como lo hará el agente de pautas
    chk = openpyxl.load_workbook(salida_p, data_only=True).worksheets[0]
    print("Brief generado:", salida_p)
    for c in ["B1", "B2", "B3", "B4", "B5", "B6", "B7", "B8", "K2", "K3", "K4",
              "K5", "K6", "K7", "K8", "K9", "K11", "K12", "K13"]:
        print(f"  {c}: {chk[c].value}")
    for i in range(len(anuncios)):
        f = FILAS[0] + i
        print(f"  Anuncio {i+1}: {chk[f'B{f}'].value} · {chk[f'C{f}'].value} · {chk[f'H{f}'].value}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
